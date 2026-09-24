"""First-pass workbook profile, following the outside-in review order: structure before formulas.

It records what a reviewer needs before opening the file: sheets and their visibility, used
ranges, tables, defined names, formula density, external links, data validation and macros.
It never modifies the workbook.
"""
import logging
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from .limits import LIGHT_ABOVE, expanded_size, too_big

log = logging.getLogger("groundwork.worker")


def profile_workbook(path: Path) -> dict:
    refused = too_big(path)
    if refused:
        return {"type": "workbook", "summary": refused, "review_flags": [refused]}
    if expanded_size(path) > LIGHT_ABOVE:
        return profile_workbook_light(path)
    wb = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    sheets = []
    total_formulas = 0
    for ws in wb.worksheets:
        formulas = cells = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                cells += 1
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
        total_formulas += formulas
        header = [str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1), []) if c.value is not None][:30]
        sheets.append({
            "name": ws.title, "state": ws.sheet_state, "dimensions": ws.dimensions,
            "max_row": ws.max_row, "max_column": ws.max_column, "filled_cells": cells, "formulas": formulas,
            "tables": list(ws.tables.keys()), "merged_ranges": len(ws.merged_cells.ranges),
            "hidden_rows": sum(1 for d in ws.row_dimensions.values() if d.hidden),
            "hidden_columns": sum(1 for d in ws.column_dimensions.values() if d.hidden),
            "data_validations": len(ws.data_validations.dataValidation),
            "conditional_formats": len(ws.conditional_formatting),
            "auto_filter": bool(ws.auto_filter.ref), "first_row": header,
        })
    names = []
    try:
        names = [{"name": n, "refers_to": str(dn.attr_text)} for n, dn in wb.defined_names.items()][:200]
    except Exception:  # openpyxl versions differ here; names are a nice-to-have
        log.info("couldn't read defined names in %s", path.name)
    with zipfile.ZipFile(path) as z:
        members = z.namelist()
    has_macros = any(m.endswith("vbaProject.bin") for m in members)
    external = [m for m in members if m.startswith("xl/externalLinks/") and m.endswith(".xml")]
    queries = any(m.startswith("customXml/") or "connections.xml" in m for m in members)
    hidden = [s["name"] for s in sheets if s["state"] != "visible"]
    flags = []
    if hidden:
        flags.append(f"{len(hidden)} hidden sheet(s): {', '.join(hidden)}")
    if external:
        flags.append(f"links to {len(external)} other workbook(s)")
    if has_macros:
        flags.append("contains macros")
    if queries:
        flags.append("has data connections or queries")
    summary = (f"Workbook with {len(sheets)} sheet(s), {total_formulas} formula cells"
               + (f"; {'; '.join(flags)}" if flags else "") + ".")
    return {"type": "workbook", "summary": summary, "sheets": sheets, "defined_names": names,
            "external_links": len(external), "has_macros": has_macros, "has_connections": queries,
            "review_flags": flags}


def profile_workbook_light(path: Path) -> dict:
    """For very large workbooks: stream each sheet once in read-only mode. Counts sheets, sizes and
    formulas, but not hidden rows, merged cells, validations or conditional formats."""
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    sheets, total = [], 0
    try:
        for ws in wb.worksheets:
            formulas = cells = 0
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is None:
                        continue
                    cells += 1
                    if isinstance(v, str) and v.startswith("="):
                        formulas += 1
            total += formulas
            sheets.append({"name": ws.title, "state": ws.sheet_state, "dimensions": ws.calculate_dimension(),
                           "filled_cells": cells, "formulas": formulas})
    finally:
        wb.close()
    with zipfile.ZipFile(path) as z:
        members = z.namelist()
    has_macros = any(m.endswith("vbaProject.bin") for m in members)
    external = [m for m in members if m.startswith("xl/externalLinks/") and m.endswith(".xml")]
    hidden = [s["name"] for s in sheets if s["state"] != "visible"]
    flags = ["very large workbook: read in light mode, so hidden rows, merged cells and validations weren't checked"]
    if hidden:
        flags.append(f"{len(hidden)} hidden sheet(s): {', '.join(hidden)}")
    if external:
        flags.append(f"links to {len(external)} other workbook(s)")
    if has_macros:
        flags.append("contains macros")
    return {"type": "workbook", "summary": f"Workbook with {len(sheets)} sheet(s), {total} formula cells; " + "; ".join(flags) + ".",
            "sheets": sheets, "external_links": len(external), "has_macros": has_macros, "review_flags": flags}
