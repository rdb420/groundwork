import io
import json
from urllib.parse import unquote

from openpyxl import Workbook

from app import worker
from app.ai import generate as gen
from tests.conftest import SENT, sign_in

H = {"x-requested-with": "groundwork"}


def xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Arrears"
    ws.append(["Tenant", "Room", "Owing"])
    ws.append(["A", 1, 120])
    ws.append(["B", 2, 80])
    ws["C4"] = "=SUM(C2:C3)"
    hidden = wb.create_sheet("Lookups")
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_domain_refused_quietly(client):
    r = client.post("/api/auth/request", json={"email": "someone@gmail.com"})
    assert r.status_code == 200 and "someone@gmail.com" not in SENT


def test_token_single_use(client):
    token = sign_in(client, "staff@example.com.au")
    r = client.post("/api/auth/verify", json={"token": unquote(token)})
    assert r.status_code == 400


def test_mutation_needs_header(client):
    sign_in(client, "staff@example.com.au")
    assert client.put("/api/auth/me", json={"display_name": "Sam"}).status_code == 403
    assert client.put("/api/auth/me", json={"display_name": "Sam", "team": "Tenancy"}, headers=H).status_code == 200
    assert client.get("/api/auth/me").json()["display_name"] == "Sam"


def test_upload_profile_and_visibility(client):
    sign_in(client, "staff@example.com.au")
    procs = client.get("/api/processes").json()
    arrears = next(p for p in procs if p["name"] == "Rent collection and arrears")
    meta = {"title": "Weekly arrears tracker", "kind": "spreadsheet", "layer": "workaround",
            "frequency": "weekly", "personal_info": "yes", "process_ids": [arrears["id"]],
            "new_process_names": ["Payment plan follow-up"]}
    r = client.post("/api/artifacts", headers=H, data={"meta": json.dumps(meta)},
                    files={"file": ("arrears.xlsx", xlsx_bytes())})
    assert r.status_code == 200, r.text
    aid = r.json()["id"]
    while worker.run_once():
        pass
    a = client.get(f"/api/artifacts/{aid}").json()
    assert a["status"] == "processed"
    assert "hidden sheet" in a["profile"]["summary"]
    assert a["profile"]["sheets"][0]["formulas"] == 1
    assert {p["name"] for p in a["processes"]} == {"Rent collection and arrears", "Payment plan follow-up"}

    sign_in(client, "outsider@example.com.au")
    assert client.get("/api/artifacts").json() == []
    assert client.get(f"/api/artifacts/{aid}").status_code == 404

    sign_in(client, "lead@example.com.au")
    assert aid in {a["id"] for a in client.get("/api/artifacts").json()}


def test_rejects_unknown_extension(client):
    sign_in(client, "staff@example.com.au")
    r = client.post("/api/artifacts", headers=H, data={"meta": json.dumps({"title": "x"})},
                    files={"file": ("run.exe", b"MZ")})
    assert r.status_code == 415


def test_board_save_conflict_and_generation(client, monkeypatch):
    sign_in(client, "lead@example.com.au")
    procs = client.get("/api/processes").json()
    arrears = next(p for p in procs if p["name"] == "Rent collection and arrears")
    b = client.post("/api/boards", headers=H, json={"title": "Arrears as it runs today",
                                                    "process_id": arrears["id"]}).json()
    doc = {"nodes": [
        {"id": "lane1", "type": "lane", "position": {"x": 0, "y": 0}, "width": 800, "height": 200,
         "data": {"label": "Property manager"}},
        {"id": "s", "type": "bpmnStart", "position": {"x": 20, "y": 80}, "data": {"label": "Rent due"}},
        {"id": "t", "type": "bpmnTask", "position": {"x": 150, "y": 70}, "data": {"label": "Check bank feed"}},
        {"id": "g", "type": "bpmnGateway", "position": {"x": 350, "y": 70}, "data": {"label": "Paid?"}},
        {"id": "n", "type": "sticky", "position": {"x": 150, "y": 20}, "data": {"label": "Done in Excel"}},
    ], "edges": [{"id": "e1", "source": "s", "target": "t"}, {"id": "e2", "source": "t", "target": "g"}],
        "viewport": None}
    r = client.put(f"/api/boards/{b['id']}", headers=H, json={"version": b["version"], "doc": doc})
    assert r.status_code == 200 and r.json()["version"] == 2
    stale = client.put(f"/api/boards/{b['id']}", headers=H, json={"version": 1, "doc": doc})
    assert stale.status_code == 409

    # The linked spreadsheet contains personal info and the provider is cloud: refuse.
    r = client.post(f"/api/boards/{b['id']}/generate", headers=H, json={"mode": "sop"})
    assert r.status_code == 409 and "personal information" in r.json()["detail"]

    captured = {}

    def fake_complete(system, user, json_mode=True):
        captured["prompt"] = user
        return json.dumps({"markdown": "# Arrears\n[TO CONFIRM: threshold]", "proposal": {"nodes": [], "edges": []}}), \
            "ollama", "test-model"

    monkeypatch.setattr(gen, "complete", fake_complete)
    monkeypatch.setattr(gen, "is_local", lambda: True)
    r = client.post(f"/api/boards/{b['id']}/generate", headers=H, json={"mode": "sop"})
    assert r.status_code == 200, r.text
    assert "Property manager" in captured["prompt"] and "Done in Excel" in captured["prompt"]
    assert "needs at least two paths out" in captured["prompt"]
    assert "Weekly arrears tracker" in captured["prompt"]
    assert r.json()["markdown"].startswith("# Arrears")
    assert client.post(f"/api/drafts/{r.json()['id']}/accept", headers=H).status_code == 200


def test_recording_requires_consent(client):
    sign_in(client, "lead@example.com.au")
    b = client.post("/api/boards", headers=H, json={"title": "Maintenance"}).json()
    assert client.post(f"/api/boards/{b['id']}/recordings", headers=H, json={"consent_note": ""}).status_code == 422
    rec = client.post(f"/api/boards/{b['id']}/recordings", headers=H,
                      json={"consent_note": "Sam and Priya agreed at 10:02"}).json()
    r = client.post(f"/api/recordings/{rec['id']}/chunks?seq=0", headers=H, files={"file": ("c.webm", b"\x1a\x45")})
    assert r.status_code == 200
    assert client.post(f"/api/recordings/{rec['id']}/end", headers=H).status_code == 200
    assert len(client.get(f"/api/boards/{b['id']}/transcript").json()) == 1
